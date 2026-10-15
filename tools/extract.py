"""Local document-to-table/text extraction. Never executes input instructions or formulas."""
import argparse
import base64
import csv
import datetime
import hashlib
import io
import json
import pathlib
import sys
import zipfile

MAX_BYTES = 20 * 1024 * 1024


def extract(filename, payload):
    if len(payload) > MAX_BYTES:
        raise ValueError("File exceeds 20 MB")
    name = pathlib.Path(filename).name
    suffix = pathlib.Path(name).suffix.lower()
    source = {"file": name, "sha256": hashlib.sha256(payload).hexdigest()}
    result = {"source": source, "tables": [], "documents": [], "warnings": []}
    if suffix in (".xlsx", ".docx"):
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded Office document exceeds 100 MB")
    if suffix == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, (datetime.datetime, datetime.date)):
                        value = value.isoformat()[:10]
                    if cell.data_type == "f":
                        result["warnings"].append(f"Formula retained as data, not evaluated: {sheet.title}!{cell.coordinate}")
                    values.append(value if value is not None else "")
                rows.append(values)
                if len(rows) > 50000:
                    raise ValueError("Worksheet exceeds 50,000 rows")
            result["tables"].append({"name": sheet.title, "rows": rows, "source": {**source, "sheet": sheet.title}})
        book.close()
    elif suffix == ".pdf":
        from pypdf import PdfReader
        book = PdfReader(io.BytesIO(payload))
        if len(book.pages) > 200:
            raise ValueError("PDF exceeds 200 pages")
        for page_number, page in enumerate(book.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                result["warnings"].append(f"Page {page_number}: no extractable text; OCR/manual review required")
            result["documents"].append({"id": f"{name}:p{page_number}", "name": name, "text": text, "page": page_number, "source": {**source, "page": page_number}})
        result["warnings"].append("PDF text is searchable evidence, not an automatically accepted validation record")
    elif suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(payload))
        lines = [p.text for p in document.paragraphs]
        result["documents"].append({"id": f"{name}:paragraphs", "name": name, "text": "\n".join(lines), "line": 1, "source": {**source, "section": "paragraphs"}})
        for index, table in enumerate(document.tables, 1):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            result["tables"].append({"name": f"Table {index}", "rows": rows, "source": {**source, "table": index}})
            result["documents"].append({"id": f"{name}:table{index}", "name": name, "text": "\n".join(" | ".join(row) for row in rows), "line": 1, "source": {**source, "table": index}})
        result["warnings"].append("DOCX paragraphs use paragraph-line citations, not fabricated page numbers")
    elif suffix in (".txt", ".md"):
        result["documents"].append({"id": name, "name": name, "text": payload.decode("utf-8-sig"), "line": 1, "source": source})
    elif suffix == ".csv":
        result["tables"].append({"name": pathlib.Path(name).stem, "rows": list(csv.reader(io.StringIO(payload.decode("utf-8-sig")))), "source": source})
    elif suffix == ".json":
        result["bundle"] = json.loads(payload.decode("utf-8-sig"))
    else:
        raise ValueError("Supported formats: XLSX, PDF, DOCX, CSV, JSON, TXT, MD")
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("file", nargs="?")
    parser.add_argument("--stdin-json", action="store_true")
    parser.add_argument("--out")
    args = parser.parse_args()
    if args.stdin_json:
        request = json.loads(sys.stdin.read(MAX_BYTES * 2))
        filename, payload = request["filename"], base64.b64decode(request["contentBase64"], validate=True)
    elif args.file:
        target = pathlib.Path(args.file)
        filename, payload = target.name, target.read_bytes()
    else:
        parser.error("Supply a file or --stdin-json")
    value = json.dumps(extract(filename, payload), ensure_ascii=False, default=str)
    if args.out:
        pathlib.Path(args.out).write_text(value, encoding="utf-8")
    else:
        sys.stdout.write(value)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        main()
    except Exception as exc:
        sys.stderr.write(f"Extraction failed: {exc}\n")
        sys.exit(2)
